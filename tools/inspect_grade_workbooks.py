"""Read-only inventory of the 18 user-supplied grade manuscript workbooks."""
import json
import re
from pathlib import Path
from collections import Counter
import openpyxl
from lxml import html

SOURCE=Path(r'C:\Users\1992k\Desktop\홈페이지 정리\참고자료\원고모음(엑셀)')
ROOT=Path(__file__).resolve().parents[1]
GRADES=['초3','초4','초5','초6','중1','중2','중3','고1','고2']

def filename(grade,subject):
    return f'{grade} {subject}학원'+('' if grade=='초3' and subject=='영어' else ' 원고')+'.xlsx'

def inspect():
    reports=[]
    parents=json.loads((ROOT/'tools/data/subject-pages/manifest.json').read_text(encoding='utf-8'))['records']
    expected=[p['locality'] for p in parents if p['subject']=='영어']
    for grade in GRADES:
        for subject in ['영어','수학']:
            path=SOURCE/filename(grade,subject)
            book=openpyxl.load_workbook(path,read_only=True,data_only=True)
            sheets=[]
            for sheet in book:
                cells=[];anomalies=[];mappings=[]
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None and cell.row > len(expected):
                            continue
                        if not isinstance(cell.value,str) or len(cell.value)<300:
                            anomalies.append({'cell':cell.coordinate,'value':str(cell.value)[:300]});continue
                        raw=cell.value.replace('_x000D_','\r')
                        if '<h1' in raw:
                            doc=html.fromstring(raw)
                            title=' '.join(doc.xpath('string(//h1)').split())
                            headings=doc.xpath('//h2/text()')[:5]
                            intro=' '.join(doc.xpath('string(//p[contains(@class,"intro")])').split())
                        else:
                            title=raw.splitlines()[0][:130];headings=[];intro=raw[:200]
                        cells.append({'cell':cell.coordinate,'title':title,'length':len(raw),'intro':intro,'headings':headings})
                        locality=expected[cell.row-1]
                        compact=lambda x:re.sub(r'\s+','',x)
                        if compact(locality) not in compact(title):
                            mappings.append({'cell':cell.coordinate,'expected':locality,'title':title,
                                             'foundInBody':compact(locality) in compact(raw)})
                sheets.append({'name':sheet.title,'rows':sheet.max_row,'cols':sheet.max_column,'manuscripts':len(cells),
                               'uniqueTitles':len(set(c['title'] for c in cells)), 'anomalies':anomalies,'mappingChecks':mappings})
            reports.append({'file':path.name,'sheets':sheets})
            print(json.dumps({'file':path.name,'sheets':[{'name':s['name'],'manuscripts':s['manuscripts'],'anomalies':s['anomalies'][:10],'mappingChecks':len(s['mappingChecks'])} for s in sheets]},ensure_ascii=False))
            book.close()
    target=ROOT/'tools/data/grade-pages/workbook-inspection.json'
    target.parent.mkdir(exist_ok=True,parents=True)
    target.write_text(json.dumps(reports,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')

if __name__=='__main__':inspect()
