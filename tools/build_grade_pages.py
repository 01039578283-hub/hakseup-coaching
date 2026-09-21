"""Expand existing subject hubs with nine grade guides from 18 read-only workbooks.

No deployment side effects. Public copying, factual claims and course availability
are bounded to the verified directory; workbook HTML is never executed or reused.
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from urllib.parse import unquote, urljoin
import openpyxl
from lxml import etree, html
import build_subject_pages as subject
from grade_editorial import GRADES, PROFILES, grade_name, selected_checks, practice_check
from inspect_grade_workbooks import filename

base = subject.base
ROOT, DAY, SOURCE = subject.ROOT, subject.DAY, subject.SOURCE
DATA = ROOT/'tools/data/grade-pages'
e, url, section, buttons = base.e, base.url, base.section, base.buttons
READ = subject.read_json
EXPECTED_ERRORS = {('중3','영어',97), *(('고2','수학',r) for r in [140,210,240,274,304,310])}
# Row layout is identical across all 18 books and the paired 371-row subject books.
# These rows contain wrong/alternate neighborhood words, not a change of center.
REVIEWED_ROW_REPAIRS = {('초3','영어',243),('초3','영어',322),('고1','영어',151),('고2','수학',115)}
ALIASES = {'경산사동':['사동'],'강릉교동':['교동'],'화성태안':['태안읍','태안동'],
           '수완지구':['수완동'],'봉담2지구':['봉담3지구'],'신불당':['불당동']}


def digest(text): return hashlib.sha256(text.encode('utf-8')).hexdigest()
def course_for(c, subject_name): return next(x for x in c['courses'] if x['subject']==subject_name)
def status_for(c, r):
    course=course_for(c,r['subject'])
    return 'confirmed' if r['grade'] in course['grades'] else 'pending' if r['grade'] in course['pendingGrades'] else 'unlisted'


def availability(c, r):
    course=course_for(c,r['subject']); grade=r['grade']; name=r['subject']; state=status_for(c,r)
    if state=='confirmed':
        return f'{c["routeName"]}의 {name} 안내 학년에 {grade} 과정이 포함되어 있습니다. 현재 모집 요일과 세부 수업 범위는 상담에서 확인하세요.'
    if state=='pending':
        return f'{c["routeName"]}의 {grade} {name}는 개설 조건 확인이 필요한 항목입니다. 아래 내용은 학습 방법 안내이며 등록 가능한 수업을 확정한 안내는 아닙니다.'
    return f'확인된 {c["routeName"]} {name} 개설 학년에는 {grade} 과정이 표시되어 있지 않습니다. 아래 학습 방법을 참고하되 수업 가능 여부는 지점에 먼저 확인하세요.'


def extract():
    parents=READ(subject.DATA/'manifest.json')['records']
    indexed={s:[r for r in parents if r['subject']==s] for s in ['영어','수학']}
    assert all(len(v)==371 for v in indexed.values())
    assert [r['locality'] for r in indexed['영어']]==[r['locality'] for r in indexed['수학']]
    records=[];sources=[];repairs=[]
    for grade in GRADES:
        for name in ['영어','수학']:
            path=SOURCE/filename(grade,name)
            book=openpyxl.load_workbook(path,read_only=True,data_only=True)
            assert book.sheetnames==['Sheet1'], path
            sheet=book['Sheet1']
            # Never shift row association when a workbook contains #ERROR!.
            values=list(sheet.iter_rows(values_only=True))
            assert all(not any(row) for row in values[371:]), (path,'Unexpected extra manuscript')
            valid=0
            for i,parent in enumerate(indexed[name],1):
                value=values[i-1][0]
                replacement=(grade,name,i) in EXPECTED_ERRORS
                if replacement:
                    assert value=='#ERROR!',(path,i,'Expected error cell changed; review again')
                    raw=parent['sourceText']; source_title=''; source_paragraphs=[]
                    mapping='New grade guidance authored from the same-neighborhood subject guide and grade profile; source cell is #ERROR!.'
                    repairs.append({'file':path.name,'cell':'A'+str(i),'locality':parent['locality'],'reason':mapping})
                else:
                    assert isinstance(value,str) and len(value)>300,(path,i,value)
                    doc=html.fromstring(value.replace('_x000D_','\r'))
                    for n in doc.xpath('//script|//style|//iframe'): n.drop_tree()
                    source_title=subject.norm(doc.xpath('string(//h1)'))
                    source_paragraphs=[subject.norm(p.text_content()) for p in doc.xpath('//p') if len(subject.norm(p.text_content()))>=50]
                    raw='\n'.join(source_paragraphs)
                    flat=subject.compact(doc.text_content()); loc=parent['locality']
                    evidence=[loc,loc.split()[-1],*ALIASES.get(loc,[])]
                    if not any(subject.compact(v) in flat for v in evidence):
                        assert (grade,name,i) in REVIEWED_ROW_REPAIRS,(path,i,loc,source_title)
                    mapping='Same physical row in paired subject manifest; title/body locality evidence checked.'
                    if subject.compact(loc) not in subject.compact(source_title) or (grade,name,i) in REVIEWED_ROW_REPAIRS:
                        mapping='Reviewed alternate/wrong locality wording against verified subject manifest and adjacent workbook rows; source addresses and branch claims not imported.'
                        repairs.append({'file':path.name,'cell':'A'+str(i),'locality':loc,'sourceTitle':source_title,'reason':mapping})
                    valid+=1
                checks,scores=selected_checks(name,raw)
                records.append({
                    'locality':parent['locality'],'region':parent['region'],'branch':parent['branch'],
                    'subject':name,'grade':grade,'title':parent['locality']+' '+grade+' '+name+'학원',
                    'branchPath':parent['parentPath'],'parentPath':parent['path'],'path':parent['path']+grade+'/',
                    'focuses':[p['key'] for p in checks],'focusScores':scores,
                    'sourceFile':path.name,'sourceSheet':'Sheet1','sourceCell':'A'+str(i),'sourceTitle':source_title,
                    'sourceTextSha256':digest(raw),'sourceParagraphHashes':[digest(p) for p in source_paragraphs],
                    'replacementForError':replacement,'mappingEvidence':mapping,
                })
            sources.append({'file':path.name,'sha256':base.sha(path),'sheet':'Sheet1','expectedRows':371,'validManuscripts':valid,'replacementRows':371-valid})
            book.close()
    assert len(records)==6678 and len({r['path'] for r in records})==6678
    old=READ(DATA/'manifest.json') if (DATA/'manifest.json').exists() else {}
    manifest={'createdAt':old.get('createdAt',datetime.now(timezone.utc).isoformat()),'modifiedAt':DAY,
              'sources':sources,'repairs':repairs,'records':records}
    base.save(DATA/'manifest.json',manifest)
    print(base.j({'workbooks':18,'newGradePages':len(records),'sourceErrorsReplaced':sum(r['replacementForError'] for r in records),'reviewedMappings':len(repairs)}))
    return manifest


def source_checks(r):
    from subject_editorial import ENGLISH,MATH
    pool=ENGLISH if r['subject']=='영어' else MATH
    return [next(p for p in pool if p['key']==key) for key in r['focuses']]


def render(c, r, lookup, center_graph):
    grade, name, locality, path = r['grade'],r['subject'],r['locality'],r['path']
    p=PROFILES[(grade,name)]; checks=source_checks(r); course=course_for(c,name); state=status_for(c,r)
    title=r['title']+' | '+p['title']
    lead=f'{locality} {grade} {name}학원을 알아볼 때 먼저 점검할 내용은 {p["summary"]}입니다. '+availability(c,r)
    description=f'{r["title"]} 학습 안내. {p["summary"]}, {checks[0]["summary"]}, '+(f'{c["routeName"]}의 개설 조건과 상담 준비를 확인하세요.' if state=='confirmed' else f'{c["routeName"]}의 학년 개설 여부를 먼저 확인하세요.')
    body=base.hero(c['region']+' · '+c['district']+' / '+c['routeName']+' / '+grade_name(grade),e(r['title']),lead,
                   [('학년별 공부 방법','#grade-focus'),('지점·개설 조건','#grade-center'),('다른 학년 보기',r['parentPath']+'#grade-pages')])
    body+='<article class="hs-article hg-article" aria-label="'+e(r['title'])+' 학습 안내">'
    body+=section('grade-focus',p['title'],f'<p>{e(p["explanation"])}</p><ol class="hc-steps">'+''.join(f'<li><strong>{e(a)}</strong><p>{e(b)}</p></li>' for a,b in p['steps'])+'</ol>','GRADE LEARNING')
    body+=section('practice-example','직접 해볼 짧은 연습',
                   '<div class="hg-example"><p class="hg-example-label">연습 문제</p><p>'+e(p['example'])+'</p><h3>확인할 내용</h3><p>'+e(p['answer'])+'</p></div>'+
                   '<p class="hc-small">학습 방법을 이해하기 위한 예시입니다. 실제 학생의 성적 사례나 지점의 수업 실적이 아니며, 학교·교재에서 현재 배운 범위에 맞춰 활용하세요.</p>','TRY IT YOURSELF')
    body+=section('focused-checks',grade+' '+name+'에서 함께 점검할 두 가지',
                   '<div class="hg-checks">'+''.join('<section class="hs-practice"><h3>'+e(x['title'])+'</h3><p>'+e(x['explain'])+'</p><p>'+e(practice_check(x,p))+'</p><p class="hc-small">남길 기록: '+e(x['record'])+'</p></section>' for x in checks)+'</div>','FOCUS & FEEDBACK')
    body+=section('review-plan','복습이 끝났는지 확인하는 기준',f'<p>{e(p["review"])}</p><p>{e(p["next_step"])}</p><div class="hg-record"><h3>공책에 남길 세 줄</h3><ol><li>도움 없이 해 본 첫 답과 판단</li><li>달라진 부분과 고친 이유</li><li>답을 가리고 다시 확인한 결과</li></ol></div>'+buttons([('오답 기록을 복습에 연결하기','/학습가이드/#wrong-answer')]),'REVIEW CHECK')
    body+='</article>'
    level={'초':'초등','중':'중등','고':'고등'}[grade[0]]
    schools=c['schools'].get(level,[])
    school_text=' · '.join(schools)
    course_status=availability(c,r)
    body+=section('grade-center',c['routeName']+' 방문과 '+grade+' '+name+' 상담 준비',
                   f'<p class="hg-status hg-status-{state}"><strong>개설 학년 확인</strong> {e(course_status)}</p>'+
                   f'<dl class="hc-facts"><div><dt>실제 학원</dt><dd>{e(c["displayName"])}</dd></div><div><dt>방문 주소</dt><dd>{e(c["address"])}</dd></div><div><dt>{e(name)} 안내 범위</dt><dd>{e(course["label"])}</dd></div><div><dt>준비할 자료</dt><dd>{e(p["materials"])}</dd></div></dl>'+
                   ''.join(f'<p class="hc-notice">{e(n)}</p>' for n in course['notes'])+
                   (f'<p>상담 참고 학교: {e(school_text)}. 학교명은 자료를 준비할 때 참고하는 목록이며 제휴나 학교별 전용반을 뜻하지 않습니다.</p>' if schools else '<p>재학 학교와 현재 배우는 단원·교재를 함께 알려주세요. 학교별 진도나 평가 방식은 실제 학교 자료로 확인합니다.</p>')+
                   f'<p class="hc-small">{e(locality)}은 {e(c["routeName"])}의 상담 참고 지역입니다. 별도의 {e(locality)} 지점을 뜻하지 않으며, 실제 통학 동선은 위 주소로 확인하세요.</p>'+
                   buttons([('지점 전체 정보와 학습 공간',r['branchPath']),('상담에 준비할 질문','/진단상담/')]),'CENTER & CONSULTATION')
    faqs=[(p['question'],p['explanation']),
          (checks[0]['question'],checks[0]['answer']),
          (c['routeName']+'에서 '+grade+' '+name+' 수업을 바로 신청할 수 있나요?',course_status+' '+(' '.join(course['notes'])+' ' if course['notes'] else '')+'최종 교육비, 시간표와 모집 상태는 상담에서 확인하세요.'),
          (grade+' '+name+' 상담에 무엇을 준비하면 좋나요?',p['materials']+'를 준비하세요. '+p['review'])]
    body+=base.faq(faqs)
    other_name='수학' if name=='영어' else '영어'
    related=[(locality+' '+name+'학원 전체 학년',r['parentPath']+'#grade-pages'),(c['routeName']+' 기본정보',r['branchPath'])]
    sibling=lookup[(locality,other_name,grade)]
    related.append((sibling['title'],sibling['path']))
    position=GRADES.index(grade)
    for neighbor in [position-1,position+1]:
        if 0<=neighbor<len(GRADES):
            child=lookup[(locality,name,GRADES[neighbor])]
            related.append((child['title'],child['path']))
    body+=section('related-pages','이어서 볼 학년·과목 안내',buttons(related),'RELATED LEARNING')
    body+=subject.subject_media(c,r['title'])
    org=copy.deepcopy(next(n for n in center_graph if n.get('@id')==url(r['branchPath'])+'#center'))
    # Do not imply a course is offered for an unlisted or pending grade.
    service=None
    if state=='confirmed':
        service={'@type':'Service','@id':url(path)+'#service','name':c['displayName']+' '+grade+' '+name+' 수업 안내',
                 'serviceType':name+' 학습코칭','provider':{'@id':org['@id']},
                 'audience':{'@type':'EducationalAudience','educationalRole':'student','audienceType':grade_name(grade)+' 학생'},
                 'areaServed':{'@type':'AdministrativeArea','name':c['region']+' '+c['district']},
                 'description':course_status+' '+' '.join(course['notes'])}
    extra=[org,*([service] if service else []),base.itemlist(path,'이어서 볼 학년·과목 안내',related)]
    base.page(path,title,description,body,[('지점안내','/지점안내/'),(c['region'],'/지점안내/'+c['region']+'/'),
               (c['routeName'],r['branchPath']),(locality+' '+name+'학원',r['parentPath']),(grade,path)],
               c['primaryMedia']['representative']['src'],extra,article=True)
    doc=html.parse(str(subject.file_for(path))).getroot();_,graph=subject.graph_of(doc)
    web=next(n for n in graph if n.get('@type')=='WebPage'); article=next(n for n in graph if n.get('@type')=='Article')
    about=[{'@id':org['@id']},{'@type':'Thing','name':grade_name(grade)+' '+name+' 공부 방법'}]
    web.update({'isPartOf':{'@id':url(r['parentPath'])+'#webpage'},'mainEntity':{'@id':article['@id']},'about':about,'datePublished':DAY})
    article.update({'abstract':lead,'about':about,'isPartOf':{'@id':web['@id']},'datePublished':DAY,
                    'audience':{'@type':'EducationalAudience','audienceType':grade_name(grade)+' 학생·학부모'},
                    'mentions':[{'@type':'EducationalOrganization','name':s} for s in schools]})
    if service: web['mentions']=[{'@id':service['@id']}]
    save_page(path,doc,graph)
    return {'title':title,'description':description,'status':state}


def save_page(path,doc,graph):
    if not doc.xpath('//link[contains(@href,"branch-grades.css")]'):
        etree.SubElement(doc.find('head'),'link',rel='stylesheet',href='/assets/branch-grades.css?v=20260921')
    subject.save_document(path,doc,graph)


def enrich_parents(parents, records, centers):
    grouped=defaultdict(list)
    for r in records:grouped[r['parentPath']].append(r)
    by_center={base.route(c):c for c in centers}
    for parent in parents:
        path=parent['path']; c=by_center[parent['parentPath']]; children=grouped[path]
        assert len(children)==9
        children.sort(key=lambda r:GRADES.index(r['grade']))
        doc=html.parse(str(subject.file_for(path))).getroot();_,graph=subject.graph_of(doc)
        markup='<p>현재 학년에서 해볼 연습과 복습 기준을 골라 보세요. 학습 안내와 실제 개설 여부를 구분해 각 글에 표시했습니다.</p><div class="hg-grade-groups">'
        for prefix,label in [('초','초등학교'),('중','중학교'),('고','고등학교')]:
            markup+='<section><h3>'+label+'</h3><div class="hg-grade-links">'
            for r in children:
                if r['grade'][0]!=prefix:continue
                label='학습 방법 · 개설 학년 포함' if status_for(c,r)=='confirmed' else '학습 방법 · 개설 확인 필요'
                markup+='<a href="'+e(r['path'])+'"><strong>'+e(r['grade']+' '+r['subject'])+'</strong><span>'+e(label)+'</span></a>'
            markup+='</div></section>'
        markup+='</div>'
        subject.replace_section(doc,'grade-pages',section('grade-pages','학년별 '+parent['subject']+' 학습 안내',markup,'CHOOSE YOUR GRADE'),after='faq')
        links=doc.xpath('//section[contains(@class,"hc-hero")]//div[@class="hc-links"]')[0]
        if not links.xpath('a[@href="#grade-pages"]'):
            links.append(subject.html_fragment('<a href="#grade-pages">학년별 공부 방법 <span aria-hidden="true">↗</span></a>'))
        ident=url(path)+'#grade-list'
        graph=[n for n in graph if n.get('@id')!=ident]
        node=base.itemlist(path,'학년별 '+parent['subject']+' 학습 안내',[(r['title'],r['path']) for r in children]);node['@id']=ident;graph.append(node)
        web=next(n for n in graph if n.get('@id')==url(path)+'#webpage')
        childids={url(r['path'])+'#webpage' for r in children}
        haspart=[n for n in web.get('hasPart',[]) if n.get('@id') not in childids and n.get('@id')!=url(path)+'#grade-pages']
        web['hasPart']=haspart+[{'@type':'WebPageElement','@id':url(path)+'#grade-pages','name':'학년별 '+parent['subject']+' 학습 안내'}]+[{'@id':url(r['path'])+'#webpage'} for r in children]
        web['dateModified']=DAY
        article=next(n for n in graph if n.get('@type')=='Article')
        article['dateModified']=DAY
        heading='학년별 '+parent['subject']+' 학습 안내'
        if heading not in article['articleSection']:article['articleSection'].append(heading)
        save_page(path,doc,graph)


def discover(manifest, info):
    records=manifest['records']; release=READ(base.DATA/'release-pages.json')
    for path,details in info.items():release[path]={k:v for k,v in details.items() if k!='status'}
    base.save(base.DATA/'release-pages.json',release)
    sm=etree.parse(str(ROOT/'sitemap.xml'));root=sm.getroot();ns={'s':'http://www.sitemaps.org/schemas/sitemap/0.9'}
    old={unquote(n.findtext('s:loc',namespaces=ns)):n for n in root}
    for r in records:
        u=url(r['path']);n=old.get(unquote(u))
        if n is None:
            n=etree.SubElement(root,'{'+ns['s']+'}url');etree.SubElement(n,'{'+ns['s']+'}loc').text=u
        date=n.find('s:lastmod',ns)
        if date is None:date=etree.SubElement(n,'{'+ns['s']+'}lastmod')
        date.text=DAY
    base.write(ROOT/'sitemap.xml',etree.tostring(sm,encoding='unicode',pretty_print=True))
    # Fixed pre-grade feed snapshot keeps regeneration stable; no fabricated dates.
    feed=DATA/'rss-before-grades.xml'
    if not feed.exists():base.write(feed,(ROOT/'rss.xml').read_text(encoding='utf-8'))
    rss=etree.parse(str(feed));channel=rss.find('channel');previous=list(channel.findall('item'))
    for n in previous:channel.remove(n)
    selected=[]
    regions=[v for v in base.REGIONS if any(r['region']==v for r in records)]
    for i in range(30):
        region=regions[i%len(regions)]; grade=GRADES[i%len(GRADES)]; name=['영어','수학'][i%2]
        candidate=next(r for r in records if r['region']==region and r['grade']==grade and r['subject']==name)
        if candidate not in selected:selected.append(candidate)
    created=datetime.fromisoformat(manifest['createdAt']).astimezone(timezone.utc).strftime('%a, %d %b %Y %H:%M:%S GMT')
    for r in selected:
        item=etree.SubElement(channel,'item');detail=info[r['path']]
        for k,v in [('title',detail['title']),('link',url(r['path'])),('guid',url(r['path'])),('description',detail['description']),('pubDate',created)]:etree.SubElement(item,k).text=v
        doc=html.parse(str(subject.file_for(r['path'])));main=doc.xpath('//main')[0]
        for el in main.xpath('.//*[@href or @src or @srcset]'):
            for attr in ['href','src']:
                if el.get(attr):el.set(attr,urljoin(url(r['path']),el.get(attr)))
            if el.get('srcset'):el.set('srcset',', '.join(urljoin(url(r['path']),part.strip().split()[0])+' '+' '.join(part.strip().split()[1:]) for part in el.get('srcset').split(',')))
        etree.SubElement(item,'{http://purl.org/rss/1.0/modules/content/}encoded').text=etree.CDATA(html.tostring(main,encoding='unicode').replace('</source>',''))
    for item in previous[:50-len(selected)]:channel.append(item)
    build=channel.find('lastBuildDate')
    if build is not None:build.text=created
    base.write(ROOT/'rss.xml',etree.tostring(rss,encoding='unicode',pretty_print=True))
    llms=ROOT/'llms.txt';text=llms.read_text(encoding='utf-8');text=re.sub(r'\n## 동네별 과목과 학년 학습 안내[\s\S]*','',text)
    text+='\n## 동네별 과목과 학년 학습 안내\n\n지점안내 → 지역 → 실제 지점 → 동네별 영어학원·수학학원 → 학년 순서로 탐색합니다. 초3·초4·초5·초6·중1·중2·중3·고1·고2 안내에는 학습 예시와 복습 기준이 있습니다. 학습 안내의 대상과 실제 지점의 개설 학년은 다를 수 있으며 각 글 상단과 개설 조건을 확인해야 합니다. 원본 엑셀의 홍보 문구를 지점의 성과·일정으로 간주하지 않습니다.\n\n- [지점에서 학습 안내 찾기]('+url('/지점안내/')+')\n- [전체 공개 페이지 사이트맵]('+base.DOMAIN+'/sitemap.xml)\n'
    base.write(llms,text)
    base.save(DATA/'release-pages.json',{'modifiedAt':DAY,'newArticles':[r['path'] for r in records],
               'subjectHubs':list(dict.fromkeys(r['parentPath'] for r in records)),'rssNewItems':[r['path'] for r in selected],
               'availabilityCounts':dict(Counter(v['status'] for v in info.values()))})


def update_from_manifest():
    manifest=READ(DATA/'manifest.json');records=manifest['records'];centers=READ(base.DATA/'centers.json')['centers']
    parents=READ(subject.DATA/'manifest.json')['records'];by_center={base.route(c):c for c in centers}
    lookup={(r['locality'],r['subject'],r['grade']):r for r in records}
    graphs={p:subject.graph_of(html.parse(str(subject.file_for(p))))[1] for p in by_center}
    info={}
    for i,r in enumerate(records,1):
        info[r['path']]=render(by_center[r['branchPath']],r,lookup,graphs[r['branchPath']])
        if i%1000==0:print(base.j({'renderedGradePages':i}),flush=True)
    enrich_parents(parents,records,centers)
    discover(manifest,info)
    print(base.j({'newGradePages':len(records),'subjectHubs':len(parents),'availabilityCounts':dict(Counter(v['status'] for v in info.values())),'deploy':False}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--extract-only',action='store_true');parser.add_argument('--reuse-manifest',action='store_true');args=parser.parse_args()
    if not args.reuse_manifest:extract()
    if not args.extract_only:update_from_manifest()
